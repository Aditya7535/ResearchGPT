"""
DOCX text extractor using python-docx.

Preserves heading style names so the chunker can use structural
information (Heading 1, Heading 2, etc.) in addition to regex matching.
"""

from __future__ import annotations

import logging
from pathlib import Path

from app.services.ingestion.extractors.base import ExtractedDocument

logger = logging.getLogger(__name__)


def extract_docx(file_path: str | Path) -> ExtractedDocument:
    """
    Extract text and heading structure from a DOCX file.

    Parameters
    ----------
    file_path:
        Path to the ``.docx`` file.

    Returns
    -------
    ExtractedDocument
        ``raw_text``   – full plain text joined by newlines.
        ``paragraphs`` – list of ``(style_name, text)`` tuples.
                         Style names are Word styles, e.g.
                         ``"Heading 1"``, ``"Normal"``, ``"List Paragraph"``.

    Raises
    ------
    ImportError
        If ``python-docx`` is not installed.
    FileNotFoundError
        If the file does not exist.
    """
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for DOCX extraction. "
            "Install with: pip install python-docx"
        ) from exc

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    logger.info("Extracting DOCX: %s", path.name)

    doc = Document(str(path))
    paragraphs: list[tuple[str, str]] = []
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else "Normal"
        paragraphs.append((style, text))
        lines.append(text)

    raw_text = "\n".join(lines)
    word_count = len(raw_text.split())

    logger.info(
        "DOCX extracted: %d paragraphs, ~%d words — %s",
        len(paragraphs), word_count, path.name,
    )

    return ExtractedDocument(
        filename=path.name,
        file_type="docx",
        raw_text=raw_text,
        paragraphs=paragraphs,
        metadata={
            "paragraph_count": len(paragraphs),
            "word_count": word_count,
            "source_path": str(path),
        },
    )
